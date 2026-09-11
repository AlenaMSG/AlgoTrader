"""
PredictionTracker: guarda historial de análisis por símbolo en Analisis_Predicción/.

Estructura de carpetas:
  Analisis_Predicción/
    AAPL/
      2022/  AAPL_Real.xlsx        ← datos de mercado del año 2022
      2023/  AAPL_Real.xlsx        ← datos de mercado del año 2023
      ...
      2026/  AAPL_Real.xlsx        ← datos de mercado del año 2026
             AAPL_Prediccion.xlsx  ← predicciones hechas en 2026

_Real.xlsx    → año del dato de mercado  (fácil de cargar por rango de años para reentrenar)
_Prediccion.xlsx → año del scan          (agrupa predicciones por el año en que se hicieron)

Formato Excel:
  - Cabecera: texto blanco negrita sobre fondo azul oscuro, nombres en MAYÚSCULAS
  - Columna SEÑAL: solo color de letra (verde / rojo / naranja / gris), sin fondo
"""
from __future__ import annotations

from datetime import datetime, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from src.analysis.feature_engineer import FEATURE_COLS

OHLCV_COLS = ["open", "high", "low", "close", "volume"]

REAL_FIELDS = ["fecha"] + OHLCV_COLS + FEATURE_COLS + ["retorno_real_siguiente"]

PRED_FIELDS = [
    "fecha_scan",
    "train_desde",
    "train_hasta",
    "fecha_pred",
    "señal",
    "pred_return_pct",
    "da_pct",
    "precio_scan",
    "pct_real",
    "fecha_real",
    "acerto",
    "pct_aciertos",
]

# ── Estilos ───────────────────────────────────────────────────────────────────
_HDR_FONT  = Font(bold=True, color="FFFFFF", size=10)
_HDR_FILL  = PatternFill(fill_type="solid", fgColor="1A2A5A")
_HDR_ALIGN = Alignment(horizontal="center", vertical="center")
_NO_FILL   = PatternFill(fill_type=None)

_SIGNAL_FONTS = {
    "comprar_fuerte": Font(bold=True, color="00CC88"),
    "vender_fuerte":  Font(bold=True, color="FF4455"),
    "debil":          Font(bold=True, color="FFAA00"),
    "neutral":        Font(bold=True, color="888888"),
}


def _signal_style(signal: str) -> str:
    s = signal.upper()
    if "COMPRAR" in s and "DÉBIL" not in s and "DEBIL" not in s:
        return "comprar_fuerte"
    if "VENDER" in s and "DÉBIL" not in s and "DEBIL" not in s:
        return "vender_fuerte"
    if "DÉBIL" in s or "DEBIL" in s:
        return "debil"
    return "neutral"


def _apply_header(ws) -> None:
    for cell in ws[1]:
        cell.font      = _HDR_FONT
        cell.fill      = _HDR_FILL
        cell.alignment = _HDR_ALIGN


def _apply_signal_col(ws, col_idx: int) -> None:
    for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
        for cell in row:
            if cell.value:
                cell.font = _SIGNAL_FONTS[_signal_style(str(cell.value))]
                cell.fill = _NO_FILL


def _autofit(ws) -> None:
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 30)


def _write_xlsx(df: pd.DataFrame, path: Path, signal_col: str | None = None) -> None:
    df_out = df.rename(columns=str.upper)
    df_out.to_excel(path, index=False, engine="openpyxl")
    wb = load_workbook(path)
    ws = wb.active
    _apply_header(ws)
    if signal_col and signal_col.upper() in df_out.columns:
        _apply_signal_col(ws, df_out.columns.get_loc(signal_col.upper()) + 1)
    _autofit(ws)
    wb.save(path)


def _read_xlsx(path: Path, usecols: list[str] | None = None) -> pd.DataFrame:
    """Lee un xlsx y normaliza las cabeceras a minúsculas (se guardan en MAYÚSCULAS)."""
    df = pd.read_excel(path, engine="openpyxl", dtype=str).fillna("")
    df.columns = [str(c).lower() for c in df.columns]
    if usecols is not None:
        df = df[list(usecols)]
    return df


# ──────────────────────────────────────────────────────────────────────────────
# Helpers de rutas
# ──────────────────────────────────────────────────────────────────────────────

def _real_path(base: Path, symbol: str, year: int) -> Path:
    d = base / symbol / str(year)
    d.mkdir(parents=True, exist_ok=True)
    return d / f"{symbol}_Real.xlsx"


def _pred_path(base: Path, symbol: str, year: int) -> Path:
    d = base / symbol / str(year)
    d.mkdir(parents=True, exist_ok=True)
    return d / f"{symbol}_Prediccion.xlsx"


def _next_trading_day(dt: datetime) -> str:
    nxt = dt + timedelta(days=1)
    while nxt.weekday() >= 5:
        nxt += timedelta(days=1)
    return nxt.strftime("%Y-%m-%d")


# ──────────────────────────────────────────────────────────────────────────────
# % Aciertos: acumulado corriente dentro del fichero
# ──────────────────────────────────────────────────────────────────────────────

def _recalc_pct_aciertos(df: pd.DataFrame) -> pd.DataFrame:
    """
    Recalcula pct_aciertos como porcentaje acumulado corriente para cada fila
    que tenga acerto = SI o NO (en orden de aparición en el fichero).
    """
    df = df.copy()
    si_total = 0
    filled   = 0
    for idx in df.index:
        val = str(df.at[idx, "acerto"]).strip().upper()
        if val in ("SI", "NO"):
            filled += 1
            if val == "SI":
                si_total += 1
            df.at[idx, "pct_aciertos"] = f"{si_total / filled * 100:.1f}%"
        else:
            df.at[idx, "pct_aciertos"] = ""
    return df


# ──────────────────────────────────────────────────────────────────────────────
# API pública
# ──────────────────────────────────────────────────────────────────────────────

def save_real_data(base_dir: Path, symbol: str, df_feat: pd.DataFrame) -> None:
    """
    Guarda/actualiza {symbol}_Real.xlsx organizado por año del dato de mercado.
    Solo añade filas nuevas (dedup por fecha).
    """
    df = df_feat.copy()
    df["fecha"] = df.index.strftime("%Y-%m-%d")
    df["retorno_real_siguiente"] = (df["close"].shift(-1) / df["close"] - 1) * 100
    df["_year"] = df.index.year

    cols = ["fecha"] + [c for c in OHLCV_COLS + FEATURE_COLS + ["retorno_real_siguiente"]
                        if c in df.columns]

    for year, group in df.groupby("_year"):
        out_path = _real_path(base_dir, symbol, int(year))
        out = group[cols].reset_index(drop=True)

        if out_path.exists():
            existing = set(_read_xlsx(out_path, usecols=["fecha"])["fecha"])
            out = out[~out["fecha"].isin(existing)]
            if out.empty:
                continue
            wb = load_workbook(out_path)
            ws = wb.active
            for _, row in out.iterrows():
                ws.append([row[c] if c in out.columns else "" for c in cols])
            _autofit(ws)
            wb.save(out_path)
        else:
            _write_xlsx(out, out_path)


def save_prediction(
    base_dir:    Path,
    symbol:      str,
    result:      dict,
    train_desde: str,
    train_hasta: str,
) -> bool:
    """
    Añade una fila a {symbol}_Prediccion.xlsx del año actual.
    Devuelve False si ya existe una predicción para hoy (dedup por fecha_scan).
    """
    now        = datetime.now()
    fecha_scan = now.strftime("%Y-%m-%d")
    fecha_pred = _next_trading_day(now)
    out_path   = _pred_path(base_dir, symbol, now.year)

    if out_path.exists():
        df_ex = _read_xlsx(out_path, usecols=["fecha_scan"])
        if fecha_scan in df_ex["fecha_scan"].values:
            return False

    metrics  = result.get("metrics", {})
    quote    = result.get("quote", {})
    pred_pct = result["pred_return"] * 100

    row = {
        "fecha_scan":      fecha_scan,
        "train_desde":     train_desde,
        "train_hasta":     train_hasta,
        "fecha_pred":      fecha_pred,
        "señal":           result["signal"],
        "pred_return_pct": f"{pred_pct:+.2f}",
        "da_pct":          f"{metrics.get('oos_dir_acc', 0) * 100:.1f}",
        "precio_scan":     f"{quote.get('ask_price', 0):.4f}",
        "pct_real":        "",
        "fecha_real":      "",
        "acerto":          "",
        "pct_aciertos":    "",
    }

    if out_path.exists():
        wb  = load_workbook(out_path)
        ws  = wb.active
        ws.append([row[f] for f in PRED_FIELDS])
        sig_col = PRED_FIELDS.index("señal") + 1
        cell = ws.cell(row=ws.max_row, column=sig_col)
        cell.font = _SIGNAL_FONTS[_signal_style(row["señal"])]
        cell.fill = _NO_FILL
        _autofit(ws)
        wb.save(out_path)
    else:
        df_new = pd.DataFrame([row], columns=PRED_FIELDS)
        _write_xlsx(df_new, out_path, signal_col="señal")

    return True


def update_actuals(base_dir: Path, symbols: list[str] | None = None) -> int:
    """
    Para cada {symbol}_Prediccion.xlsx (todos los años) donde fecha_pred ya pasó
    y pct_real está vacío, descarga el cierre real con yfinance y rellena
    pct_real, fecha_real, acerto y recalcula pct_aciertos.

    También actualiza retorno_real_siguiente en _Real.xlsx del año correspondiente.
    Retorna el número de filas completadas.
    """
    import yfinance as yf

    today   = datetime.now().date()
    updated = 0

    if symbols:
        sym_dirs = [base_dir / s for s in symbols if (base_dir / s).is_dir()]
    else:
        sym_dirs = [p for p in base_dir.iterdir() if p.is_dir()] if base_dir.exists() else []

    for sym_dir in sym_dirs:
        symbol = sym_dir.name

        # Recorrer todas las carpetas de año
        year_dirs = sorted([p for p in sym_dir.iterdir() if p.is_dir() and p.name.isdigit()])
        for year_dir in year_dirs:
            pred_xlsx = year_dir / f"{symbol}_Prediccion.xlsx"
            if not pred_xlsx.exists():
                continue

            df = _read_xlsx(pred_xlsx)
            # Asegurar columna pct_aciertos
            if "pct_aciertos" not in df.columns:
                df["pct_aciertos"] = ""

            mask = (df["pct_real"] == "") & (
                pd.to_datetime(df["fecha_pred"], errors="coerce").dt.date < today
            )
            pending = df.index[mask].tolist()
            if not pending:
                continue

            changed = False
            for idx in pending:
                try:
                    scan_dt     = pd.Timestamp(df.at[idx, "fecha_scan"])
                    precio_scan = float(df.at[idx, "precio_scan"])
                    end_dt      = scan_dt.date() + timedelta(days=14)

                    hist = yf.Ticker(symbol).history(
                        start=str(scan_dt.date()), end=str(end_dt),
                        interval="1d", auto_adjust=True,
                    )
                    if len(hist) < 2:
                        continue

                    c1         = float(hist["Close"].iloc[1])
                    actual_ret = (c1 - precio_scan) / precio_scan * 100
                    fecha_real = hist.index[1].strftime("%Y-%m-%d")
                    pred_pct   = float(df.at[idx, "pred_return_pct"].replace("+", "") or 0)
                    acerto     = "SI" if (pred_pct > 0) == (actual_ret > 0) else "NO"

                    df.at[idx, "pct_real"]   = f"{actual_ret:+.2f}"
                    df.at[idx, "fecha_real"] = fecha_real
                    df.at[idx, "acerto"]     = acerto
                    changed  = True
                    updated += 1

                    # Actualizar retorno_real_siguiente en _Real.xlsx del año del scan
                    real_xlsx = year_dir / f"{symbol}_Real.xlsx"
                    if real_xlsx.exists():
                        rdf    = _read_xlsx(real_xlsx)
                        mask_r = rdf["fecha"] == str(scan_dt.date())
                        if mask_r.any():
                            rdf.loc[mask_r, "retorno_real_siguiente"] = f"{actual_ret:+.2f}"
                            _write_xlsx(rdf, real_xlsx)

                except Exception:
                    pass

            if changed:
                df = _recalc_pct_aciertos(df)
                _write_xlsx(df, pred_xlsx, signal_col="señal")

    return updated
